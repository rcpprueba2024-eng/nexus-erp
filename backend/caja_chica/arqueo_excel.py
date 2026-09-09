"""Llena el arqueo de caja en el MISMO Excel que ya usa RCP a mano (ver
caja_chica/plantillas_excel/arqueo_caja.xlsx — plantilla real que trajo el
jefe, con sus fórmulas, colores y formato originales intactos).

Regla de oro: acá solo se escriben las celdas que en la plantilla real son
datos escritos a mano cada día. TODO lo demás (TOTAL A, TOTAL B, B-A,
TICKET PROM., sumas por vendedor, etc.) son fórmulas que ya trae la
plantilla — si esas fórmulas se llegaran a sobrescribir con un número fijo
calculado aparte en Python, se corre el riesgo real de que algún día se
desalineen del cálculo que el negocio ya conoce y confía. Por eso nunca se
tocan.
"""
import os
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "plantillas_excel", "arqueo_caja.xlsx")

FILA_CREDITOS_INICIO = 5
FILA_INGRESOS_INICIO = 15
FILA_GASTOS_INICIO = 15
FILA_VENDEDOR_VENTAS_INICIO = 35
FILA_VENDEDOR_EQUIPOS_INICIO = 43


def _set(ws, coord, value):
    cell = ws[coord]
    if isinstance(cell, MergedCell):
        return
    cell.value = value


def arqueo_caja_workbook(arqueo):
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb["Arqueo"]
    ws_efectivo = wb["efectivo"]

    ws.title = arqueo.fecha.strftime("%-d-%-m-%Y") if os.name != "nt" else arqueo.fecha.strftime("%#d-%#m-%Y")

    _set(ws, "O3", arqueo.fecha)
    # F4 es la única celda de la plantilla para esto y está en dólares — el
    # monto en córdobas (guardado aparte) se convierte acá con el tipo de
    # cambio del día, no se deja como una fila/fórmula nueva en el Excel.
    tipo_cambio = float(arqueo.tipo_cambio) if arqueo.tipo_cambio else 1
    caja_inicial_usd = float(arqueo.caja_chica_inicial or 0) + (float(arqueo.caja_chica_inicial_cordobas or 0) / tipo_cambio if tipo_cambio else 0)
    _set(ws, "F4", caja_inicial_usd)

    # Créditos pendientes (snapshot del día) — hasta 7 filas
    for i, c in enumerate((arqueo.creditos or [])[:7]):
        fila = FILA_CREDITOS_INICIO + i
        _set(ws, f"H{fila}", c.get("cliente", ""))
        _set(ws, f"I{fila}", c.get("factura") or None)
        _set(ws, f"J{fila}", c.get("ot", ""))
        _set(ws, f"K{fila}", c.get("monto") or 0)
        _set(ws, f"L{fila}", c.get("vence") or None)

    # Columna B — lo realmente cerrado por canal (efectivo y gastos los
    # calcula la fórmula sola, ver más abajo)
    _set(ws, "P5", float(arqueo.pagos_cheque_cierre))
    _set(ws, "P6", float(arqueo.pagos_transferencia_cierre))
    _set(ws, "P7", float(arqueo.compras_credex_cierre))
    _set(ws, "P8", float(arqueo.cierre_pos))
    _set(ws, "P9", float(arqueo.creditos_cierre))

    # (X) Ingresos por ventas del día — hasta 15 filas.
    # "monto" (Q(U$)) se deja realmente en blanco (no 0) cuando no viene
    # dado: el TICKET PROM. por vendedor (K30/L30) usa AVERAGEIF, que trata
    # blanco distinto de 0 — un 0 escrito ahí sí cuenta en el promedio y
    # daría un ticket promedio equivocado.
    for i, it in enumerate((arqueo.ingresos_ventas or [])[:15]):
        fila = FILA_INGRESOS_INICIO + i
        _set(ws, f"D{fila}", it.get("descripcion", ""))
        _set(ws, f"F{fila}", it.get("monto") if it.get("monto") not in (None, "") else None)
        _set(ws, f"G{fila}", it.get("tipo") or None)
        _set(ws, f"H{fila}", it.get("banco") or None)
        _set(ws, f"I{fila}", it.get("factura") or None)
        _set(ws, f"J{fila}", it.get("ots") or None)
        _set(ws, f"K{fila}", it.get("vendedor") or None)
        _set(ws, f"L{fila}", it.get("empresa") or None)

    # (Y) Gastos de caja chica — hasta 10 filas
    for i, g in enumerate((arqueo.gastos_caja_chica or [])[:10]):
        fila = FILA_GASTOS_INICIO + i
        _set(ws, f"N{fila}", g.get("descripcion", ""))
        _set(ws, f"P{fila}", g.get("monto") or 0)

    _set(ws, "N27", arqueo.recibos_caja_utilizados or None)
    _set(ws, "N29", arqueo.comentarios or None)

    # Detalle de ventas del día — columna ANTERIOR (el ACUM. lo suma la fórmula)
    _set(ws, "G33", float(arqueo.ventas_dia_anterior))
    _set(ws, "G34", float(arqueo.total_com_anterior))
    _set(ws, "G35", float(arqueo.total_cca_anterior))
    _set(ws, "G36", float(arqueo.repuestos_uso_interno_anterior))
    _set(ws, "G37", float(arqueo.rep_acc_insumo_externa_anterior))
    _set(ws, "G38", float(arqueo.computadoras_anterior))

    # Detalle de ventas por vendedor — 4 filas fijas (J35:J38). "Venta del
    # día" y "Acumulado MES" los calcula la fórmula sola contra la tabla de
    # ingresos por nombre de vendedor — por eso el nombre acá debe
    # escribirse EXACTAMENTE igual al que se usó en esa tabla.
    for i in range(4):
        fila = FILA_VENDEDOR_VENTAS_INICIO + i
        it = (arqueo.ventas_por_vendedor or [])[i] if i < len(arqueo.ventas_por_vendedor or []) else None
        _set(ws, f"J{fila}", (it or {}).get("vendedor") or None)
        _set(ws, f"L{fila}", (it or {}).get("acumulado_semanal") or 0)
        _set(ws, f"N{fila}", (it or {}).get("semana_anterior") or 0)
    _set(ws, "S35", float(arqueo.meta_semanal))
    _set(ws, "S37", float(arqueo.meta_mensual))

    # Equipos recibidos por día, por vendedor — 3 filas fijas (D43:D45)
    for i in range(3):
        fila = FILA_VENDEDOR_EQUIPOS_INICIO + i
        it = (arqueo.equipos_por_vendedor or [])[i] if i < len(arqueo.equipos_por_vendedor or []) else None
        _set(ws, f"D{fila}", (it or {}).get("vendedor") or None)
        _set(ws, f"E{fila}", (it or {}).get("clientes_atendidos") or 0)
        _set(ws, f"G{fila}", (it or {}).get("com") or 0)
        _set(ws, f"H{fila}", (it or {}).get("cca") or 0)
        _set(ws, f"I{fila}", (it or {}).get("repuestos") or 0)

    # Proyección de cierre
    _set(ws, "O42", arqueo.dia_del_mes)
    _set(ws, "O44", arqueo.dias_mes)
    _set(ws, "L46", float(arqueo.tipo_cambio) if arqueo.tipo_cambio is not None else None)

    # Hoja "efectivo" — conteo físico de billetes/monedas (la fórmula del
    # arqueo principal suma esto solo, ver O45/O46 en la hoja "Arqueo")
    denom_usd = [
        (5, arqueo.billetes_usd_100), (6, arqueo.billetes_usd_50), (7, arqueo.billetes_usd_20),
        (8, arqueo.billetes_usd_10), (9, arqueo.billetes_usd_5), (10, arqueo.billetes_usd_1),
    ]
    for fila, cantidad in denom_usd:
        _set(ws_efectivo, f"B{fila}", cantidad)

    denom_nio = [
        (5, arqueo.billetes_nio_1000), (6, arqueo.billetes_nio_500), (7, arqueo.billetes_nio_200),
        (8, arqueo.billetes_nio_100), (9, arqueo.billetes_nio_50), (10, arqueo.billetes_nio_20),
        (11, arqueo.billetes_nio_10), (12, arqueo.billetes_nio_5), (13, arqueo.billetes_nio_1),
    ]
    for fila, cantidad in denom_nio:
        _set(ws_efectivo, f"F{fila}", cantidad)

    return wb


def arqueo_caja_excel_bytes(arqueo):
    wb = arqueo_caja_workbook(arqueo)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
