import os
from datetime import datetime
from io import BytesIO

from django.conf import settings
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.image import Image as XlImage
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PRIMARY_HEX = "1D4ED8"
HEADER_FILL = PatternFill(start_color=PRIMARY_HEX, end_color=PRIMARY_HEX, fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10.5)
TITLE_FONT = Font(bold=True, size=15, color="0F172A")
SUBTITLE_FONT = Font(size=9.5, color="64748B", italic=True)
ZEBRA_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
TOTALS_FILL = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
TOTALS_FONT = Font(bold=True, size=10)
THIN_BORDER = Border(bottom=Side(style="thin", color="E2E8F0"))
MONEY_FORMAT = '"$"#,##0.00'
# Fila completa a resaltar cuando el registro necesita atención (ej. un
# empleado con alguna falta en el periodo) — mismo rojo en todo el sistema.
DANGER_FILL = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")
DANGER_FONT = Font(color="FFFFFF", bold=True, size=10)

LOGO_PATH = os.path.join(settings.MEDIA_ROOT, "branding", "rcp-logo.png")


def _insertar_encabezado(ws, title, n_cols):
    """Logo (si existe) + título + fecha de generación, en las primeras 3 filas
    (fila 4 es donde arranca la tabla — ver `header_row` en build_sheet)."""
    logo_col_offset = 0
    if os.path.exists(LOGO_PATH):
        img = XlImage(LOGO_PATH)
        img.width, img.height = 40, 40
        ws.add_image(img, "A1")
        logo_col_offset = 1

    # Fila 1: título (con altura suficiente para el logo). Fila 2: subtítulo
    # con la fecha, con su propia altura normal — antes esta fila se
    # comprimía a 6px y el texto quedaba montado sobre el encabezado de la
    # tabla. Fila 3 queda como separador en blanco antes de la tabla (fila 4).
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 16

    title_col = get_column_letter(1 + logo_col_offset)
    last_col = max(n_cols, 3)
    ws[f"{title_col}1"] = title
    ws[f"{title_col}1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1 + logo_col_offset, end_row=1, end_column=last_col)

    generado = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws[f"{title_col}2"] = f"Generado el {generado}"
    ws[f"{title_col}2"].font = SUBTITLE_FONT
    ws.merge_cells(start_row=2, start_column=1 + logo_col_offset, end_row=2, end_column=last_col)


def build_sheet(ws, title, headers, rows, totals_row=None, money_cols=None,
                 highlight_rows=None, col_widths=None, wrap_cols=None, time_cols=None):
    """`money_cols`: columnas (1-indexadas) a formatear como moneda ($#,##0.00).
    `highlight_rows`: posiciones (0-indexadas, dentro de `rows`) que se pintan
    de rojo completas — para señalar filas que necesitan atención (ej. un
    empleado con faltas en el periodo), por encima del zebrado normal.
    `col_widths`: {columna 1-indexada: ancho} para columnas que necesitan más
    espacio que el ancho por defecto (22).
    `wrap_cols`: columnas (1-indexadas) con texto largo que debe ajustarse
    dentro de la celda en vez de desbordarse.
    `time_cols`: columnas (1-indexadas) cuyo valor numérico es una cantidad de
    HORAS (float) que debe mostrarse como duración `[h]:mm:ss` — el valor se
    convierte internamente a fracción de día, como espera Excel."""
    money_cols = set(money_cols or [])
    highlight_rows = set(highlight_rows or [])
    col_widths = col_widths or {}
    wrap_cols = set(wrap_cols or [])
    time_cols = set(time_cols or [])
    n_cols = len(headers)
    _insertar_encabezado(ws, title, n_cols)

    header_row = 4
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r, row in enumerate(rows, start=header_row + 1):
        destacada = (r - header_row - 1) in highlight_rows
        zebra = not destacada and (r - header_row) % 2 == 0
        for c, value in enumerate(row, start=1):
            if c in time_cols and isinstance(value, (int, float)):
                cell = ws.cell(row=r, column=c, value=value / 24)
                cell.number_format = "[h]:mm:ss"
            else:
                cell = ws.cell(row=r, column=c, value=value)
            cell.border = THIN_BORDER
            if destacada:
                cell.fill = DANGER_FILL
                cell.font = DANGER_FONT
            elif zebra:
                cell.fill = ZEBRA_FILL
            if c in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c in money_cols and isinstance(value, (int, float)):
                cell.number_format = MONEY_FORMAT
                cell.alignment = Alignment(horizontal="right")

    last_data_row = header_row + len(rows)
    if totals_row:
        r = last_data_row + 1
        for c, value in enumerate(totals_row, start=1):
            if c in time_cols and isinstance(value, (int, float)):
                cell = ws.cell(row=r, column=c, value=value / 24)
                cell.number_format = "[h]:mm:ss"
            else:
                cell = ws.cell(row=r, column=c, value=value)
            cell.font = TOTALS_FONT
            cell.fill = TOTALS_FILL
            if c in money_cols and isinstance(value, (int, float)):
                cell.number_format = MONEY_FORMAT
                cell.alignment = Alignment(horizontal="right")

    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{last_data_row}"
    ws.freeze_panes = f"A{header_row + 1}"

    for col in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = col_widths.get(col, 22)

    return header_row


def add_bar_chart(ws, header_row, n_rows, category_col, value_col, title, anchor=None):
    if n_rows == 0:
        return
    chart = BarChart()
    chart.title = title
    chart.y_axis.title = None
    chart.x_axis.title = None
    chart.style = 10
    # Look plano/moderno: sin el marco negro redondeado que pone openpyxl
    # por defecto, que se veía como una "caja" pesada sobre la hoja.
    chart.roundedCorners = False
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    data = Reference(ws, min_col=value_col, min_row=header_row, max_row=header_row + n_rows)
    cats = Reference(ws, min_col=category_col, min_row=header_row + 1, max_row=header_row + n_rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    for serie in chart.series:
        serie.graphicalProperties.solidFill = PRIMARY_HEX
        serie.graphicalProperties.line.noFill = True
    chart.gapWidth = 60
    chart.legend = None
    chart.width = 16
    chart.height = 8.5
    ws.add_chart(chart, anchor or f"{get_column_letter(value_col + 3)}{header_row}")


def workbook_response(wb, filename):
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def simple_workbook(title, headers, rows, totals_row=None, chart_title=None, category_col=1, value_col=None, money_cols=None):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    header_row = build_sheet(ws, title, headers, rows, totals_row, money_cols=money_cols)
    if chart_title and value_col:
        add_bar_chart(ws, header_row, len(rows), category_col, value_col, chart_title)
    return wb
