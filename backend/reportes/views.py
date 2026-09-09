from django.db.models import Sum, F, Count
from rest_framework.views import APIView
from rest_framework.response import Response
from core.permissions import role_permission
from ventas.models import Factura, DetalleFactura
from taller.models import OrdenTaller, InsumoUsado
from inventario.models import Producto
from contabilidad.models import CuentaPorCobrar
from compras.models import OrdenCompra
from rrhh.models import Empleado
from .excel_utils import simple_workbook, workbook_response, build_sheet, add_bar_chart
from .pdf_utils import simple_pdf_response, build_pdf_response
from openpyxl import Workbook


def _bucket_key(fecha, agrupacion):
    if agrupacion == "semanal":
        year, week, _ = fecha.isocalendar()
        return f"{year}-W{week:02d}"
    if agrupacion == "quincenal":
        quincena = 1 if fecha.day <= 15 else 2
        return f"{fecha.year}-{fecha.month:02d} (Q{quincena})"
    if agrupacion == "mensual":
        return f"{fecha.year}-{fecha.month:02d}"
    return fecha.isoformat()


def ventas_periodo_data(agrupacion):
    # Se suma `total_usd`, no `total`: las facturas quedan guardadas en la
    # moneda con la que pagó cada cliente (US$ o C$), así que sumar `total`
    # directo mezclaría unidades distintas en un mismo número.
    facturas = Factura.objects.exclude(estado="ANULADA").order_by("fecha")
    buckets = {}
    for f in facturas:
        key = _bucket_key(f.fecha, agrupacion)
        if key not in buckets:
            buckets[key] = {"periodo": key, "total": 0.0, "cantidad_facturas": 0}
        buckets[key]["total"] += float(f.total_usd)
        buckets[key]["cantidad_facturas"] += 1
    return sorted(buckets.values(), key=lambda x: x["periodo"])


def ventas_vendedores_data():
    qs = (
        Factura.objects.exclude(estado="ANULADA")
        .values("vendedor__id", "vendedor__username", "vendedor__first_name")
        .annotate(total_vendido=Sum("total_usd"), cantidad_facturas=Count("id"))
        .order_by("-total_vendido")
    )
    resultado = []
    for row in qs:
        total = float(row["total_vendido"] or 0)
        cantidad = row["cantidad_facturas"]
        resultado.append({
            "vendedor": row["vendedor__first_name"] or row["vendedor__username"] or "Sin asignar",
            "total_vendido": total,
            "cantidad_facturas": cantidad,
            "promedio_por_venta": round(total / cantidad, 2) if cantidad else 0,
        })
    return resultado


def categorias_vendidas_data():
    # Solo líneas de PRODUCTO con producto real (las de tipo SERVICIO o las
    # que facturan una OT ya cerrada no tienen categoría de inventario).
    qs = (
        DetalleFactura.objects.filter(tipo="PRODUCTO", producto__isnull=False)
        .exclude(factura__estado="ANULADA")
        .values("producto__categoria__id", "producto__categoria__nombre")
        .annotate(unidades_vendidas=Sum("cantidad"), total_vendido=Sum("subtotal"))
        .order_by("-total_vendido")
    )
    return [
        {
            "categoria": row["producto__categoria__nombre"] or "Sin categoría",
            "unidades_vendidas": float(row["unidades_vendidas"] or 0),
            "total_vendido": float(row["total_vendido"] or 0),
        }
        for row in qs
    ]


def gastos_proveedor_data():
    qs = (
        OrdenCompra.objects.exclude(estado="CANCELADA")
        .values("proveedor__id", "proveedor__nombre")
        .annotate(total_pagado=Sum("total"), cantidad_ordenes=Count("id"))
        .order_by("-total_pagado")
    )
    detalle = [
        {
            "proveedor": row["proveedor__nombre"],
            "total_pagado": float(row["total_pagado"] or 0),
            "cantidad_ordenes": row["cantidad_ordenes"],
        }
        for row in qs
    ]
    total = sum(d["total_pagado"] for d in detalle)
    return detalle, total


def gastos_dia_data():
    # Compras a proveedor sí tienen fecha propia; los insumos de taller no
    # (se aproximan con la fecha de ingreso de la orden a la que
    # pertenecen — es lo más cercano disponible a "cuándo se gastó").
    compras_qs = (
        OrdenCompra.objects.exclude(estado="CANCELADA")
        .values("fecha").annotate(monto=Sum("total")).order_by("fecha")
    )
    insumos_qs = (
        InsumoUsado.objects.values("orden__fecha_ingreso")
        .annotate(monto=Sum(F("cantidad") * F("producto__precio_compra")))
        .order_by("orden__fecha_ingreso")
    )
    por_dia = {}
    for row in compras_qs:
        f = row["fecha"]
        por_dia.setdefault(f, {"fecha": f, "compras_proveedores": 0.0, "insumos_taller": 0.0})
        por_dia[f]["compras_proveedores"] += float(row["monto"] or 0)
    for row in insumos_qs:
        f = row["orden__fecha_ingreso"]
        if f is None:
            continue
        por_dia.setdefault(f, {"fecha": f, "compras_proveedores": 0.0, "insumos_taller": 0.0})
        por_dia[f]["insumos_taller"] += float(row["monto"] or 0)
    data = sorted(por_dia.values(), key=lambda x: x["fecha"], reverse=True)
    for d in data:
        d["total"] = round(d["compras_proveedores"] + d["insumos_taller"], 2)
        d["fecha"] = d["fecha"].isoformat()
    return data


def estado_resultados_data():
    ventas_facturas = Factura.objects.exclude(estado="ANULADA")
    ingresos = float(ventas_facturas.aggregate(t=Sum("total_usd"))["t"] or 0)

    _, gasto_insumos = gastos_insumos_data()
    _, gasto_planilla = gastos_personal_data()
    _, gasto_compras = gastos_proveedor_data()

    gastos = [
        {"concepto": "Insumos de taller", "monto": round(gasto_insumos, 2)},
        {"concepto": "Planilla", "monto": round(gasto_planilla, 2)},
        {"concepto": "Compras a proveedores", "monto": round(gasto_compras, 2)},
    ]
    total_gastos = sum(g["monto"] for g in gastos)
    utilidad_neta = round(ingresos - total_gastos, 2)

    return {
        "ingresos": round(ingresos, 2),
        "gastos": gastos,
        "total_gastos": round(total_gastos, 2),
        "utilidad_neta": utilidad_neta,
        "margen_neto": round((utilidad_neta / ingresos * 100), 1) if ingresos else 0,
    }


def gastos_insumos_data():
    qs = (
        InsumoUsado.objects.values("producto__id", "producto__nombre")
        .annotate(cantidad_usada=Sum("cantidad"), costo_total=Sum(F("cantidad") * F("producto__precio_compra")))
        .order_by("-costo_total")
    )
    detalle = [
        {
            "producto": row["producto__nombre"],
            "cantidad_usada": row["cantidad_usada"],
            "costo_total": float(row["costo_total"] or 0),
        }
        for row in qs
    ]
    total = sum(d["costo_total"] for d in detalle)
    return detalle, total


def gastos_personal_data():
    empleados = Empleado.objects.filter(activo=True)
    detalle = [
        {"empleado": e.nombre, "puesto": ", ".join(e.puestos), "area": ", ".join(e.areas_display), "salario": float(e.salario)}
        for e in empleados
    ]
    total = sum(d["salario"] for d in detalle)
    return detalle, total


def inventario_movimiento_data():
    ventas_qs = (
        DetalleFactura.objects.filter(tipo="PRODUCTO", producto__isnull=False)
        .values("producto__id", "producto__nombre", "producto__codigo")
        .annotate(movido=Sum("cantidad"))
    )
    insumos_qs = (
        InsumoUsado.objects.values("producto__id", "producto__nombre", "producto__codigo")
        .annotate(movido=Sum("cantidad"))
    )
    movimiento = {}
    for row in list(ventas_qs) + list(insumos_qs):
        pid = row["producto__id"]
        if pid not in movimiento:
            movimiento[pid] = {"producto": row["producto__nombre"], "codigo": row["producto__codigo"], "total_movido": 0}
        movimiento[pid]["total_movido"] += row["movido"]
    for p in Producto.objects.all():
        if p.id not in movimiento:
            movimiento[p.id] = {"producto": p.nombre, "codigo": p.codigo, "total_movido": 0}
    data = list(movimiento.values())
    mas_movidos = sorted(data, key=lambda x: -x["total_movido"])[:5]
    menos_movidos = sorted(data, key=lambda x: x["total_movido"])[:5]
    return mas_movidos, menos_movidos


def cuentas_por_cobrar_data():
    cuentas = CuentaPorCobrar.objects.filter(estado="PENDIENTE").select_related("cliente")
    data = []
    for c in cuentas:
        equipo_pendiente = OrdenTaller.objects.filter(cliente=c.cliente, estado="LISTO_ENTREGA").exists()
        data.append({
            "cliente": c.cliente.nombre,
            "monto": float(c.monto),
            "fecha_vencimiento": c.fecha_vencimiento,
            "equipo_pendiente_retiro": equipo_pendiente,
        })
    return data


def taller_eficiencia_data():
    tecnicos = [e for e in Empleado.objects.all() if "TALLER" in e.areas]
    data = []
    for t in tecnicos:
        completadas_qs = OrdenTaller.objects.filter(tecnico=t, estado="ENTREGADO", fecha_entrega_real__isnull=False)
        completadas = completadas_qs.count()
        activas = OrdenTaller.objects.filter(tecnico=t).exclude(estado__in=("ENTREGADO", "CANCELADO", "LISTO_ENTREGA", "ABANDONADO")).count()
        dias_totales = sum((o.fecha_entrega_real - o.fecha_ingreso).days for o in completadas_qs)
        promedio_dias = round(dias_totales / completadas, 1) if completadas else None
        data.append({
            "tecnico": t.nombre,
            "ordenes_completadas": completadas,
            "ordenes_activas": activas,
            "promedio_dias": promedio_dias,
        })
    data.sort(key=lambda x: (-x["ordenes_completadas"], x["promedio_dias"] if x["promedio_dias"] is not None else 999))
    return data


def clientes_empresas_data():
    from datetime import date, timedelta
    from django.db.models import Q
    from ventas.models import Cliente, Empresa, DIAS_CLIENTE_FRECUENTE, UMBRAL_CLIENTE_FRECUENTE

    empresas = (
        Empresa.objects.annotate(
            n_contactos=Count("contactos", distinct=True),
            n_equipos=Count("contactos__equipos", distinct=True),
        )
        .filter(n_equipos__gt=0)
        .order_by("-n_equipos")[:15]
    )
    equipos_por_empresa = [{"empresa": e.nombre, "equipos": e.n_equipos, "contactos": e.n_contactos} for e in empresas]

    corte = date.today() - timedelta(days=DIAS_CLIENTE_FRECUENTE)
    clientes = Cliente.objects.annotate(
        n_ordenes=Count("ordenes_taller", filter=Q(ordenes_taller__fecha_ingreso__gte=corte), distinct=True),
        n_facturas=Count("facturas", filter=Q(facturas__fecha__gte=corte), distinct=True),
    ).annotate(actividad=F("n_ordenes") + F("n_facturas")).order_by("-actividad")[:15]
    top_clientes = [{"cliente": c.nombre, "actividad": c.actividad, "frecuente": c.actividad >= UMBRAL_CLIENTE_FRECUENTE} for c in clientes if c.actividad > 0]

    total_clientes = Cliente.objects.count()
    total_empresas = Empresa.objects.count()
    total_frecuentes = Cliente.objects.annotate(
        n_ordenes=Count("ordenes_taller", filter=Q(ordenes_taller__fecha_ingreso__gte=corte), distinct=True),
        n_facturas=Count("facturas", filter=Q(facturas__fecha__gte=corte), distinct=True),
    ).annotate(actividad=F("n_ordenes") + F("n_facturas")).filter(actividad__gte=UMBRAL_CLIENTE_FRECUENTE).count()

    return {
        "equipos_por_empresa": equipos_por_empresa,
        "top_clientes": top_clientes,
        "total_clientes": total_clientes,
        "total_empresas": total_empresas,
        "total_frecuentes": total_frecuentes,
    }


# ---------- Vistas JSON (frontend) ----------

class VentasPeriodoView(APIView):
    # `nunca=RRHH`: RRHH tiene "reportes" marcado para ver su propia planilla,
    # eso no debe destaparle cifras de ventas/vendedores (no es su área).
    permission_classes = [role_permission("VENTAS", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        return Response(ventas_periodo_data(request.query_params.get("agrupacion", "diario")))


class VentasPorVendedorView(APIView):
    permission_classes = [role_permission("VENTAS", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        return Response(ventas_vendedores_data())


class CategoriasVendidasView(APIView):
    permission_classes = [role_permission("VENTAS", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        return Response(categorias_vendidas_data())


class GastosInsumosView(APIView):
    permission_classes = [role_permission("TALLER", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        detalle, total = gastos_insumos_data()
        return Response({"detalle": detalle, "total_gastado": total})


class GastosPersonalView(APIView):
    # Sin `modules`: planilla es sensible y no debe abrirse por el checkbox
    # genérico "Reportes" de Configuración, solo por rol RRHH/BACKOFFICE.
    permission_classes = [role_permission("RRHH", "BACKOFFICE")]

    def get(self, request):
        detalle, total = gastos_personal_data()
        return Response({"detalle": detalle, "total_planilla": total})


class InventarioMovimientoView(APIView):
    permission_classes = [role_permission("TALLER", "BACKOFFICE", "VENTAS", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        mas, menos = inventario_movimiento_data()
        return Response({"mas_movidos": mas, "menos_movidos": menos})


class CuentasPorCobrarReporteView(APIView):
    # Sin `modules`: cuentas por cobrar es sensible, solo rol BACKOFFICE.
    permission_classes = [role_permission("BACKOFFICE")]

    def get(self, request):
        return Response(cuentas_por_cobrar_data())


class ClientesEmpresasView(APIView):
    # Mismo alcance que ClienteViewSet: Ventas/Backoffice, nunca Taller/Pasante.
    permission_classes = [role_permission("VENTAS", "BACKOFFICE", modules=("clientes", "ventas", "reportes"), nunca=("TALLER", "PASANTE", "RRHH"))]

    def get(self, request):
        return Response(clientes_empresas_data())


class TallerEficienciaView(APIView):
    permission_classes = [role_permission("TALLER", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        return Response(taller_eficiencia_data())


class GastosProveedorView(APIView):
    # Sin `modules`: cuánto se le paga a cada proveedor es sensible.
    permission_classes = [role_permission("BACKOFFICE")]

    def get(self, request):
        detalle, total = gastos_proveedor_data()
        return Response({"detalle": detalle, "total_pagado": total})


class GastosDiaView(APIView):
    permission_classes = [role_permission("BACKOFFICE")]

    def get(self, request):
        return Response(gastos_dia_data())


class EstadoResultadosView(APIView):
    permission_classes = [role_permission("BACKOFFICE")]

    def get(self, request):
        return Response(estado_resultados_data())


# ---------- Exportación a Excel y PDF ----------

class ExportarVentasPeriodoView(APIView):
    permission_classes = [role_permission("VENTAS", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        agrupacion = request.query_params.get("agrupacion", "diario")
        data = ventas_periodo_data(agrupacion)
        rows = [[d["periodo"], d["total"], d["cantidad_facturas"]] for d in data]
        wb = simple_workbook(
            f"Ventas ({agrupacion})", ["Periodo", "Total ($)", "Facturas"], rows,
            chart_title="Ventas por período", category_col=1, value_col=2, money_cols=[2],
        )
        return workbook_response(wb, f"ventas_{agrupacion}.xlsx")


class ExportarVentasPeriodoPdfView(APIView):
    permission_classes = [role_permission("VENTAS", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        agrupacion = request.query_params.get("agrupacion", "diario")
        data = ventas_periodo_data(agrupacion)
        rows = [[d["periodo"], d["total"], d["cantidad_facturas"]] for d in data]
        return simple_pdf_response(
            f"Ventas por período ({agrupacion})", ["Periodo", "Total ($)", "Facturas"], rows,
            f"ventas_{agrupacion}.pdf", money_cols=[1],
        )


class ExportarVentasVendedoresView(APIView):
    permission_classes = [role_permission("VENTAS", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        data = ventas_vendedores_data()
        rows = [[d["vendedor"], d["total_vendido"], d["cantidad_facturas"], d["promedio_por_venta"]] for d in data]
        wb = simple_workbook(
            "Vendedores", ["Vendedor", "Total vendido ($)", "Facturas", "Promedio por venta ($)"], rows,
            chart_title="Total vendido por vendedor", category_col=1, value_col=2, money_cols=[2, 4],
        )
        return workbook_response(wb, "ventas_por_vendedor.xlsx")


class ExportarVentasVendedoresPdfView(APIView):
    permission_classes = [role_permission("VENTAS", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        data = ventas_vendedores_data()
        rows = [[d["vendedor"], d["total_vendido"], d["cantidad_facturas"], d["promedio_por_venta"]] for d in data]
        return simple_pdf_response(
            "Ventas por vendedor", ["Vendedor", "Total vendido ($)", "Facturas", "Promedio por venta ($)"], rows,
            "ventas_por_vendedor.pdf", money_cols=[1, 3],
        )


class ExportarGastosProveedorView(APIView):
    permission_classes = [role_permission("BACKOFFICE")]

    def get(self, request):
        detalle, total = gastos_proveedor_data()
        rows = [[d["proveedor"], d["cantidad_ordenes"], d["total_pagado"]] for d in detalle]
        wb = simple_workbook(
            "Gastos por proveedor", ["Proveedor", "Órdenes", "Total pagado ($)"], rows,
            totals_row=["TOTAL", "", total],
            chart_title="Pagado por proveedor", category_col=1, value_col=3, money_cols=[3],
        )
        return workbook_response(wb, "gastos_por_proveedor.xlsx")


class ExportarGastosProveedorPdfView(APIView):
    permission_classes = [role_permission("BACKOFFICE")]

    def get(self, request):
        detalle, total = gastos_proveedor_data()
        rows = [[d["proveedor"], d["cantidad_ordenes"], d["total_pagado"]] for d in detalle]
        return simple_pdf_response(
            "Gastos por proveedor", ["Proveedor", "Órdenes", "Total pagado ($)"], rows,
            "gastos_por_proveedor.pdf", totals_row=["TOTAL", "", total], money_cols=[2],
        )


class ExportarGastosView(APIView):
    # Incluye planilla (salarios por empleado): sin `modules`, solo
    # RRHH/BACKOFFICE, nunca por el checkbox genérico "Reportes".
    permission_classes = [role_permission("BACKOFFICE", "RRHH")]

    def get(self, request):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Insumos de taller"
        insumos, total_insumos = gastos_insumos_data()
        rows1 = [[d["producto"], d["cantidad_usada"], d["costo_total"]] for d in insumos]
        header_row1 = build_sheet(ws1, "Gastos en insumos de taller", ["Insumo", "Cantidad usada", "Costo total ($)"], rows1,
                                   totals_row=["TOTAL", "", total_insumos], money_cols=[3])
        add_bar_chart(ws1, header_row1, len(rows1), 1, 3, "Costo por insumo")

        ws2 = wb.create_sheet("Planilla")
        personal, total_personal = gastos_personal_data()
        rows2 = [[d["empleado"], d["puesto"], d["area"], d["salario"]] for d in personal]
        build_sheet(ws2, "Gastos de personal (planilla)", ["Empleado", "Puesto", "Área", "Salario ($)"], rows2,
                    totals_row=["TOTAL PLANILLA", "", "", total_personal], money_cols=[4])

        return workbook_response(wb, "gastos.xlsx")


class ExportarGastosPdfView(APIView):
    permission_classes = [role_permission("BACKOFFICE", "RRHH")]

    def get(self, request):
        insumos, total_insumos = gastos_insumos_data()
        rows1 = [[d["producto"], d["cantidad_usada"], d["costo_total"]] for d in insumos]
        personal, total_personal = gastos_personal_data()
        rows2 = [[d["empleado"], d["puesto"], d["area"], d["salario"]] for d in personal]
        return build_pdf_response("gastos.pdf", "Gastos", [
            {
                "titulo": "Gastos en insumos de taller", "headers": ["Insumo", "Cantidad usada", "Costo total ($)"],
                "rows": rows1, "totals_row": ["TOTAL", "", total_insumos], "money_cols": [2],
            },
            {
                "titulo": "Gastos de personal (planilla)", "headers": ["Empleado", "Puesto", "Área", "Salario ($)"],
                "rows": rows2, "totals_row": ["TOTAL PLANILLA", "", "", total_personal], "money_cols": [3],
            },
        ])


class ExportarInventarioMovimientoView(APIView):
    permission_classes = [role_permission("TALLER", "BACKOFFICE", "VENTAS", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        mas, menos = inventario_movimiento_data()
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Más movidos"
        rows1 = [[d["codigo"], d["producto"], d["total_movido"]] for d in mas]
        header_row1 = build_sheet(ws1, "Productos que más se mueven", ["Código", "Producto", "Movimientos"], rows1)
        add_bar_chart(ws1, header_row1, len(rows1), 2, 3, "Más movidos")

        ws2 = wb.create_sheet("Menos movidos")
        rows2 = [[d["codigo"], d["producto"], d["total_movido"]] for d in menos]
        header_row2 = build_sheet(ws2, "Productos que menos se mueven", ["Código", "Producto", "Movimientos"], rows2)
        add_bar_chart(ws2, header_row2, len(rows2), 2, 3, "Menos movidos")

        return workbook_response(wb, "inventario_movimiento.xlsx")


class ExportarInventarioMovimientoPdfView(APIView):
    permission_classes = [role_permission("TALLER", "BACKOFFICE", "VENTAS", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        mas, menos = inventario_movimiento_data()
        rows1 = [[d["codigo"], d["producto"], d["total_movido"]] for d in mas]
        rows2 = [[d["codigo"], d["producto"], d["total_movido"]] for d in menos]
        return build_pdf_response("inventario_movimiento.pdf", "Movimiento de inventario", [
            {"titulo": "Productos que más se mueven", "headers": ["Código", "Producto", "Movimientos"], "rows": rows1},
            {"titulo": "Productos que menos se mueven", "headers": ["Código", "Producto", "Movimientos"], "rows": rows2},
        ])


class ExportarCuentasPorCobrarView(APIView):
    permission_classes = [role_permission("BACKOFFICE")]

    def get(self, request):
        data = cuentas_por_cobrar_data()
        rows = [[d["cliente"], d["monto"], str(d["fecha_vencimiento"]), "Sí" if d["equipo_pendiente_retiro"] else "No"] for d in data]
        wb = simple_workbook(
            "Cuentas por cobrar", ["Cliente", "Monto ($)", "Vencimiento", "Equipo pendiente de retiro"], rows,
            money_cols=[2],
        )
        return workbook_response(wb, "cuentas_por_cobrar.xlsx")


class ExportarCuentasPorCobrarPdfView(APIView):
    permission_classes = [role_permission("BACKOFFICE")]

    def get(self, request):
        data = cuentas_por_cobrar_data()
        rows = [[d["cliente"], d["monto"], str(d["fecha_vencimiento"]), "Sí" if d["equipo_pendiente_retiro"] else "No"] for d in data]
        return simple_pdf_response(
            "Cuentas por cobrar", ["Cliente", "Monto ($)", "Vencimiento", "Equipo pendiente de retiro"], rows,
            "cuentas_por_cobrar.pdf", money_cols=[1],
        )


class ExportarTallerEficienciaView(APIView):
    permission_classes = [role_permission("TALLER", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        data = taller_eficiencia_data()
        rows = [[d["tecnico"], d["ordenes_completadas"], d["ordenes_activas"], d["promedio_dias"] or ""] for d in data]
        wb = simple_workbook(
            "Eficiencia de técnicos", ["Técnico", "Reparadas", "Activas", "Promedio (días)"], rows,
            chart_title="Órdenes reparadas por técnico", category_col=1, value_col=2,
        )
        return workbook_response(wb, "taller_eficiencia.xlsx")


class ExportarTallerEficienciaPdfView(APIView):
    permission_classes = [role_permission("TALLER", "BACKOFFICE", modules="reportes", nunca=("RRHH",))]

    def get(self, request):
        data = taller_eficiencia_data()
        rows = [[d["tecnico"], d["ordenes_completadas"], d["ordenes_activas"], d["promedio_dias"] or ""] for d in data]
        return simple_pdf_response(
            "Eficiencia de técnicos", ["Técnico", "Reparadas", "Activas", "Promedio (días)"], rows,
            "taller_eficiencia.pdf",
        )
