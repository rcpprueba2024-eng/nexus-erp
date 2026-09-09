"""Genera el Excel de una Cotización a Cliente rellenando la plantilla real
que RCP ya usaba a mano (compras/plantillas_excel/cotizacion_cliente.xlsx,
tomada de un PPTO real del negocio) — así el documento que sale del sistema
es indistinguible del que hacían a mano, no una recreación aparte.
"""
import os
from io import BytesIO

from django.http import HttpResponse
from openpyxl import load_workbook

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "plantillas_excel", "cotizacion_cliente.xlsx")

# La tabla de ítems va de la fila 16 a la 26 (fila 27 es "SUB TOTAL", ver la
# plantilla) — se deja siempre al menos 1 fila libre al final para la nota
# de garantía, que en el formato original ocupa la fila justo después del
# último ítem (columna DESCRIPCION reutilizada como texto libre).
FILA_ITEM_INICIO = 16
FILA_ITEM_FIN = 26
MAX_ITEMS = FILA_ITEM_FIN - FILA_ITEM_INICIO


def cotizacion_cliente_workbook(cot):
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb["EN U$"]

    ws["H4"] = cot.titulo
    ws["H5"] = cot.ppto_numero
    ws["D7"] = cot.cliente_nombre
    ws["D8"] = cot.equipo_marca
    ws["D9"] = cot.servicio_producto
    ws["D10"] = cot.tipo_trabajo
    ws["D11"] = cot.periodo
    # Valor fijo, no la fórmula "=TODAY()" que traía la plantilla — así la
    # fecha de esta cotización queda congelada y no cambia si alguien vuelve
    # a abrir el archivo semanas después.
    ws["D12"] = cot.fecha
    ws["G8"] = cot.contacto_nombre
    ws["G9"] = cot.contacto_puesto
    ws["G10"] = cot.contacto_telefono
    ws["G11"] = cot.contacto_email
    ws["G12"] = cot.ruc

    tipo_cambio = float(cot.tipo_cambio)
    items = list(cot.items.all().order_by("orden", "id"))[:MAX_ITEMS]
    fila = FILA_ITEM_INICIO
    for item in items:
        precio_usd = float(item.precio_unitario_usd)
        ws[f"B{fila}"] = item.descripcion
        ws[f"F{fila}"] = item.cantidad
        # Mismo estilo que la plantilla original (ej. "=119*36.63"): una
        # fórmula, no el número ya calculado — para que en Excel se siga
        # viendo de dónde sale el precio en córdobas.
        ws[f"G{fila}"] = f"={precio_usd}*{tipo_cambio}"
        ws[f"H{fila}"] = f"=G{fila}*F{fila}"
        if item.proveedor:
            ws[f"L{fila}"] = item.proveedor
        if item.costo_proveedor_unitario is not None:
            ws[f"M{fila}"] = float(item.costo_proveedor_unitario)
            ws[f"N{fila}"] = f"=M{fila}*F{fila}"
            ws[f"O{fila}"] = f"=+H{fila}-N{fila}"
        fila += 1

    if cot.nota_garantia:
        ws[f"B{fila}"] = cot.nota_garantia

    # El total en dólares (fila 30) dividía siempre entre el tipo de cambio
    # que traía la plantilla de ejemplo (36.63) — hay que actualizar esa
    # fórmula a la tasa real de ESTA cotización si es distinta.
    ws["H30"] = f"=+H29/{tipo_cambio}"

    if cot.observaciones:
        ws["B37"] = cot.observaciones

    return wb


def cotizacion_cliente_excel_response(cot):
    wb = cotizacion_cliente_workbook(cot)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    nombre = f"PPTO_{cot.ppto_numero or cot.id}_{cot.cliente_nombre}".replace(" ", "_").replace("/", "-")
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nombre}.xlsx"'
    return response
